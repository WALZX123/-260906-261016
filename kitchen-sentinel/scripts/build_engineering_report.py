"""文件职责：中期检查说明文档的构建工具。
主要功能：组织说明正文和表格，设置 Word 样式，输出软件工程分析报告。
技术栈：Python、python-docx、WordprocessingML、表格样式。
模块联动：读取本脚本中的报告内容并写入 docs；独立于后端采集与控制，不作为实机完成的证据。
编写/修改日期（日志对应）：2026-10-14。
"""

from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "厨房安全监测平台软件工程分析与中期检查说明.docx"


def shade(cell, color):
    """为 Word 表格单元格追加底色属性，供表头和交替行样式复用，不修改业务数据。
    编写/修改日期（日志对应）：2026-10-14。
    """
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    props.append(node)


def add_table(doc, headers, rows):
    """按列标题和行数据建立 Word 表格，设置表头、字体与交替底色，再给后续正文保留间距。
    编写/修改日期（日志对应）：2026-10-14。
    """
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = str(value)
        shade(c, "1E6A5D")
        for r in c.paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            if len(t.rows) % 2 == 0:
                shade(cells[i], "F2F7F4")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.name = "Microsoft YaHei"
    doc.add_paragraph()


def p(doc, text):
    """向报告追加正文段落，统一段后间距和行距，供报告各章节复用。
    编写/修改日期（日志对应）：2026-10-14。
    """
    para = doc.add_paragraph(text)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.15


def bullets(doc, values):
    """把文本序列逐项写入 Word 项目列表，并设置条目间距。
    编写/修改日期（日志对应）：2026-10-14。
    """
    for value in values:
        item = doc.add_paragraph(value, style="List Bullet")
        item.paragraph_format.space_after = Pt(3)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)
for name, size, color in [
    ("Normal", 10, "244941"),
    ("Title", 25, "153E37"),
    ("Heading 1", 16, "1A5B4F"),
    ("Heading 2", 12, "2F7064"),
]:
    st = doc.styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("家庭厨房燃气泄漏与火灾报警系统\n软件工程分析与中期检查说明")
sub = doc.add_paragraph("Kitchen Sentinel · 本地 Web 全栈原型 · 版本 1.0.0")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p(
    doc,
    "本文以可维护、可测试、可演进的软件工程方式描述厨房安全监测平台的需求、架构和当前实现。软件仿真成果、真实硬件接入和云端部署分开记录，便于中期检查准确说明进度，并为后续生产实习或毕业设计提供开发基线。",
)
add_table(
    doc,
    ["项目属性", "当前结论"],
    [
        ("系统定位", "家庭厨房安全监测与联动处置平台，当前为本机教学仿真原型"),
        ("已实现范围", "React 浏览器端、FastAPI 服务、MySQL、设备模拟器、报警与命令闭环"),
        ("数据性质", "燃气和烟雾为 0–100 相对演示指标；温度为模拟 °C；未完成实物标定"),
        ("尚待接入", "真实传感器和执行机构、手机端、云部署、真实推送、模型服务生产化"),
        ("验收重点", "正常监测、持续异常报警、事件处理、关阀回执、历史追溯、断线与超时"),
    ],
)

doc.add_heading("1 项目背景与范围", level=1)
p(
    doc,
    "厨房安全系统同时涉及现场实时性、执行机构安全性、数据留痕和远程查看。系统采用“现场优先、平台协同”：现场控制器负责采集、阈值判断和保护动作；平台负责存储、监控、事件处理和辅助分析。网络或模型服务异常不应阻塞现场保护。",
)
doc.add_heading("1.1 当前阶段目标", level=2)
bullets(
    doc,
    [
        "用独立设备模拟器复现设备采样、心跳、本地报警、执行反馈和断线场景。",
        "实现浏览器监控工作台，查看实时状态、报警时间线、历史数据、执行命令和审计日志。",
        "使用 MySQL 持久化业务数据，用 Alembic 迁移管理表结构。",
        "为真实硬件、手机端和云端保留稳定的数据协议与业务接口。",
    ],
)
doc.add_heading("1.2 范围边界", level=2)
p(
    doc,
    "本版本不声称完成燃气浓度实物标定、真实电磁阀物理关闭验证、消防电源实际切断、手机原生应用、云平台部署或真实事故识别。模拟器执行回执只证明软件协议和状态机工作，不能替代工程安装和安全认证。",
)

doc.add_heading("2 需求分析", level=1)
add_table(
    doc,
    ["编号", "需求", "实现方式", "验收证据"],
    [
        ("FR-01", "多源采集", "接收燃气、烟雾、温度、火焰和质量标记", "模拟器连续上报，数据库存在采样"),
        ("FR-02", "本地报警", "纯规则状态机支持持续时间、回差和恢复时间", "燃气、火灾、故障场景产生不同报警"),
        ("FR-03", "报警管理", "触发、恢复、人工确认和备注分别记录", "报警详情时间线与审计记录"),
        (
            "FR-04",
            "联动控制",
            "关阀、消音、自检采用幂等键和设备回执",
            "pending → dispatched → succeeded/failed/timeout",
        ),
        ("FR-05", "历史追溯", "采样、报警、命令、审计按设备和时间关联", "曲线、明细、CSV 导出"),
        ("FR-06", "权限控制", "admin、operator、viewer 三种角色；写操作 CSRF 校验", "不同账号访问结果不同"),
        ("FR-07", "辅助分析", "规则摘要默认，模型服务可选且不改变保护动作", "详情显示 provider、证据和局限"),
        (
            "NFR",
            "可维护性与可靠性",
            "分层、事务、约束、索引、幂等、心跳和超时",
            "静态检查、接口冒烟和场景演示",
        ),
    ],
)

doc.add_heading("3 总体架构", level=1)
p(
    doc,
    "项目采用单体服务内部模块化结构。对于生产实习规模，多个微服务会增加部署和排错成本，因此先保持一个可部署进程，通过清晰模块边界为未来拆分留出空间。",
)
add_table(
    doc,
    ["层次", "组件", "职责边界"],
    [
        ("表现层", "React + TypeScript + Recharts", "页面、交互、图表、WebSocket；不产生报警结论"),
        ("接口层", "FastAPI APIRouter", "登录、设备协议、监测查询、报警、控制和设置"),
        ("业务层", "monitoring.py、rules.py", "采样接入、状态转换、报警生命周期和命令状态机"),
        ("数据层", "SQLAlchemy + Alembic + MySQL 8", "事务、约束、索引、迁移和历史数据"),
        ("边缘适配层", "simulator/main.py", "虚拟现场控制器；未来由串口、MQTT 或真实固件替换"),
        ("外部服务", "可选模型 API", "生成解释摘要；失败不影响现场规则"),
    ],
)
p(
    doc,
    "数据流：设备采样 → 身份验证 → 重复检测 → 持久化 → 规则复核 → 报警同步 → WebSocket 推送 → 页面展示。控制流：用户操作 → 权限与 CSRF → 幂等命令 → 设备拉取 → 执行 → 回执 → 状态更新 → 审计。",
)
doc.add_heading("3.1 目录与职责", level=2)
add_table(
    doc,
    ["目录", "用途"],
    [
        ("backend/api", "按业务边界组织 HTTP 路由和权限依赖"),
        ("backend/domain", "与数据库、网络无关的报警规则"),
        ("backend/services", "采样、报警、命令和快照服务"),
        ("backend/core", "配置、会话、数据库和安全工具"),
        ("simulator", "独立虚拟边缘控制器与演示场景"),
        ("frontend/src/pages", "按用户任务拆分页面"),
        ("migrations", "Alembic 数据库版本"),
        ("scripts", "初始化、启动和报告工具"),
    ],
)

doc.add_heading("4 核心领域设计", level=1)
doc.add_heading("4.1 报警状态", level=2)
p(
    doc,
    "status 表示触发条件是否仍然存在，acknowledged_at 表示人员是否已经确认，两者不能合并。无效采样不能清除已有危险报警；危险恢复后不自动重新打开燃气阀门。气体和火灾报警使用持续触发时间与回差，减少临界值抖动。",
)
add_table(
    doc,
    ["状态或字段", "含义", "处理规则"],
    [
        ("active", "条件仍存在", "保持同一事件，不为每条采样创建新事件"),
        ("recovered", "条件经过恢复确认时间后消失", "保留事件，等待人工复盘"),
        ("acknowledged_at", "人员已知悉", "不等同于条件恢复，不改变保护动作"),
        ("sensor_fault", "采样缺失或无效", "单独告警，不能把缺失数据视为正常"),
        ("offline", "超过心跳时限无上报", "显示最后读数，不冒充实时数据"),
    ],
)
doc.add_heading("4.2 控制命令状态", level=2)
p(
    doc,
    "命令使用客户端幂等键、设备行锁和过期时间。pending 表示入队，dispatched 表示设备拉取，设备回执才进入 succeeded 或 failed；超过截止时间进入 timeout，迟到成功回执不会重新改写过期命令。",
)
doc.add_heading("4.3 决策权", level=2)
bullets(
    doc,
    [
        "本地规则是即时保护的最终依据，模型只提供辅助说明。",
        "模型服务不能开阀、解除报警或改变现场保护动作。",
        "远程控制仅支持紧急关阀、消音和自检，不设计远程开阀。",
        "无物理位置反馈时，页面只显示设备回执，不宣称阀门已完成物理关闭。",
    ],
)

doc.add_heading("5 数据库设计", level=1)
add_table(
    doc,
    ["表", "关键字段", "设计目的"],
    [
        ("users / auth_sessions", "角色、密码哈希、会话哈希、CSRF", "认证与权限，不保存明文密码"),
        ("devices", "令牌哈希、心跳、latest、场景", "设备身份、在线状态和当前状态"),
        ("telemetry", "boot_id、sequence、采样和接收时间", "原始追溯；联合唯一约束防重复"),
        ("alarms", "kind、status、timeline、evidence、analysis", "报警生命周期和分析证据"),
        ("commands", "幂等键、状态、过期时间、回执", "可靠控制命令"),
        ("rule_configs", "revision、values", "可审计规则版本和并发控制"),
        ("audit_logs", "actor、action、target、detail", "关键操作留痕"),
    ],
)
p(
    doc,
    "sampled_at 与 received_at 分开保存，便于识别设备时间漂移和迟到报文。设备与接收时间、报警设备状态均建立索引，支持总览和历史查询。",
)

doc.add_heading("6 接口与协议", level=1)
add_table(
    doc,
    ["接口", "用途", "身份或权限"],
    [
        ("POST /api/v1/auth/login", "建立 HttpOnly 会话", "账号密码、登录限流"),
        ("GET /api/v1/snapshot", "读取总览快照", "已登录"),
        ("WS /api/v1/live", "推送实时快照", "会话 Cookie、来源校验"),
        ("POST /api/v1/ingest/devices/{id}/telemetry", "设备上报采样", "设备 Bearer 令牌"),
        ("GET /api/v1/ingest/devices/{id}/control", "设备拉取规则与命令", "设备令牌"),
        ("POST /api/v1/devices/{id}/commands", "创建关阀、消音、自检", "operator/admin + CSRF"),
        ("POST /api/v1/alarms/{id}/acknowledge", "确认报警并留说明", "operator/admin + CSRF"),
        ("PUT /api/v1/rules", "修改规则", "admin + revision"),
        ("GET /api/v1/telemetry/export", "导出历史 CSV", "已登录"),
    ],
)
p(
    doc,
    "设备报文包含 boot_id、sequence、sampled_at、gas、smoke、temperature、flame、valid 和 local_alarm。当前 gas、smoke 为相对指标，后续真实传感器必须补充单位、标定和质量说明。",
)

doc.add_heading("7 安全与运维设计", level=1)
bullets(
    doc,
    [
        "配置通过 .env 管理；数据库密码、会话密钥、设备令牌和模型密钥不写入源码、不返回前端。",
        "密码使用 PBKDF2-SHA256 加盐哈希；设备只保存令牌哈希。",
        "会话使用 HttpOnly、SameSite Cookie；写接口要求 CSRF；跨域来源采用白名单。",
        "控制命令使用幂等键、截止时间和设备回执；请求带 X-Request-ID 便于定位。",
        "生产部署还需增加 HTTPS、反向代理限流、密钥托管、备份、监控和真实设备安全认证。",
    ],
)

doc.add_heading("8 测试与验收计划", level=1)
add_table(
    doc,
    ["编号", "场景", "预期结果"],
    [
        ("T01", "登录与角色权限", "管理员可设置；观察员不能控制或改规则"),
        ("T02", "连续燃气异常", "形成一条燃气报警并触发模拟保护"),
        ("T03", "瞬时尖峰", "未达到持续时间，不形成持续报警"),
        ("T04", "火灾信号", "烟雾、温度或火焰触发火灾报警"),
        ("T05", "无效采样", "传感器故障，不清除已有危险报警"),
        ("T06", "重复与迟到报文", "重复不增样本；迟到样本不刷新当前状态"),
        ("T07", "关阀命令", "状态和回执可追踪，重复请求幂等"),
        ("T08", "设备离线", "出现离线报警并标明最后读数"),
        ("T09", "持久化", "重启后仍可查历史采样和报警"),
        ("T10", "模型不可用", "只影响辅助分析，不影响规则和联动"),
    ],
)
p(
    doc,
    "已完成一次接口级冒烟验证：登录返回 200；快照正常；连续 5 条高燃气相对指标样本形成活动报警；关阀命令进入 pending；前端 Vite 生产构建通过。正式提交前应按上表完成浏览器演示和截图归档。",
)

doc.add_heading("9 后续实施路线", level=1)
add_table(
    doc,
    ["阶段", "工作内容", "完成判据"],
    [
        (
            "硬件接入",
            "确定主控与传感器，完成电平、供电、引脚和低压执行器联调",
            "真实采样与执行反馈可重复，阈值有标定记录",
        ),
        ("云端部署", "迁移 MySQL、API、前端和日志，配置 HTTPS 与备份", "部署、监控和回滚可执行"),
        ("手机端", "先复用响应式 Web，必要时复用同一 API 开发原生应用", "移动权限与 Web 一致"),
        ("智能分析", "加入标注场景、规则基线、模型输出校验与成本统计", "只解释证据并表达不确定性"),
        ("工程收尾", "补齐安装、接口、测试、演示和风险清单", "第三方按文档可启动"),
    ],
)

doc.add_heading("附录 A 中期检查演示脚本", level=1)
bullets(
    doc,
    [
        "登录工作台，说明当前是本机仿真环境，燃气和烟雾是相对指标。",
        "选择“正常监测”，展示设备在线、实时曲线和执行器状态。",
        "切换“燃气持续升高”，展示曲线上升、报警创建和模拟本地保护。",
        "进入报警中心，展示编号、证据、时间线，填写处理说明并生成规则摘要。",
        "进入联动控制发送紧急关阀，说明命令下发与设备回执是两个阶段。",
        "查询历史并导出 CSV，最后演示上报中断或传感器故障。",
    ],
)
doc.add_heading("附录 B 现场说明口径", level=1)
p(
    doc,
    "当前中期成果重点是软件端和系统集成基础：本地 Web 工作台、后端接口、数据库模型、设备协议、报警状态机、模拟器和测试闭环已经完成。真实硬件、手机端和云端沿用当前 API 与数据模型继续扩展。",
)

for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
