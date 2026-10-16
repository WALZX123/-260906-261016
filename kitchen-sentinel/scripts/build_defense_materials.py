"""文件职责：生成可打印的录制与答辩材料。
功能：从中文指南创建Word手册与一页速查，保留标题层级和实际验收边界。
技术栈：python-docx、OOXML、Markdown基础段落；关联docs中的两份材料与交付目录。
编写/修改日期（日志对应）：2026-10-16。使用工作区文档运行环境执行，渲染另行验收。
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT.parent / 'outputs/system-defense-20261016'


def style_font(style, size):
    """统一中西文字体与字号，标题保持黑色，正文采用易读中文字体。"""
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')


def table(doc, rows):
    """创建有边框和足够留白的表格，禁止跨页拆分单行；首行重复并以深蓝配白字区分。"""
    rows = [r for r in rows if not re.match(r'^\|\s*[-:]+\s*\|', r)]
    data = [[c.strip() for c in row.strip('|').split('|')] for row in rows]
    t = doc.add_table(rows=0, cols=len(data[0]))
    t.autofit = False
    widths = ([.98, 2.75, 2.77] if len(data[0]) == 3 else [1.25, 5.25])
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    for index, values in enumerate(data):
        row = t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if index == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Inches(width)
            cell.text = value
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                node = OxmlElement('w:' + edge)
                node.set(qn('w:w'), '85')
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            tcpr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                node = OxmlElement('w:' + edge)
                node.set(qn('w:val'), 'single')
                node.set(qn('w:sz'), '4')
                node.set(qn('w:color'), 'D9D9D9')
                borders.append(node)
            tcpr.append(borders)
            if index == 0 or index % 2 == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), '24445C' if index == 0 else 'F4F6F7')
                tcpr.append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    r.font.size = Pt(11)
                    if index == 0:
                        r.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build(name, target, compact=False):
    """按源材料段落、标题和表格生成文档；手册按章节分页，速查通过紧凑段距保持一页。"""
    source = ROOT / 'docs' / name
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.66 if compact else .8)
    sec.left_margin = sec.right_margin = Inches(1)
    for label, size in [('Normal', 11), ('Title', 21 if compact else 23), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 12)]:
        style_font(doc.styles[label], size)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = 1.18 if compact else 1.32
    normal.space_after = Pt(5 if compact else 8)
    for label in ('Heading 1', 'Heading 2', 'Heading 3'):
        pf = doc.styles[label].paragraph_format
        pf.space_before, pf.space_after = Pt(12 if not compact else 8), Pt(7 if not compact else 5)
        pf.keep_with_next = True
    title = doc.styles['Title'].paragraph_format
    title.space_before, title.space_after = Pt(0), Pt(14)
    for tag in ('w:pBdr', 'w:shd'):
        node = doc.styles['Title'].element.find('.//' + tag, namespaces=doc.styles['Title'].element.nsmap)
        if node is not None:
            node.getparent().remove(node)
    header = sec.header.paragraphs[0]
    header.text = '厨房安全监测系统  功能演示与答辩材料'
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('2026年10月16日  ·  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)
    lines = source.read_text(encoding='utf-8').splitlines()
    index = 0
    pending_break = False
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line == '---PAGE---':
            pending_break = True
        elif line.startswith('|'):
            rows = [line]
            while index < len(lines) and lines[index].startswith('|'):
                rows.append(lines[index])
                index += 1
            table(doc, rows)
        elif line.startswith('# '):
            doc.add_paragraph(line[2:], style='Title')
        elif line.startswith('## '):
            heading = doc.add_paragraph(line[3:], style='Heading 1')
            if pending_break:
                heading.paragraph_format.page_break_before = True
                pending_break = False
        elif line.startswith('### '):
            doc.add_paragraph(line[4:], style='Heading 2')
        elif re.match(r'^\d+\. ', line):
            # 编号作为正文保留，避免Word继承未知编号模板，打印和导出保持一致。
            p = doc.add_paragraph(line)
            p.paragraph_format.left_indent = Inches(.06)
        elif line.startswith('- '):
            p = doc.add_paragraph('• ' + line[2:])
            p.paragraph_format.left_indent = Inches(.06)
        else:
            doc.add_paragraph(line)
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = doc.paragraphs[0].text
    doc.core_properties.subject = '实际功能演示、时序说明与答辩准备'
    OUTPUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT / target)
    (OUTPUT / name).write_bytes(source.read_bytes())
    print(str(OUTPUT / target))


if __name__ == '__main__':
    build('功能演示录制与答辩指南_2026-10-16.md', '功能演示录制与答辩指南_2026-10-16.docx')
    build('答辩上手速查_2026-10-16.md', '答辩上手速查_2026-10-16.docx', True)
